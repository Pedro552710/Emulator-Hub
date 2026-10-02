"""Rebuild catalog.json from the supplied workbook and reviewed enrichment.

Usage: python tools/import_catalog.py [workbook] [--output catalog.json]
openpyxl is only required when importing; the desktop application reads JSON.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
import sys
from urllib.parse import urlsplit

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from core.catalog import Catalog, LEGAL_NOTICE  # noqa: E402

CHECKED = "2026-10-01"
CATEGORIES = ["Nintendo", "Sony PlayStation", "Sega", "Microsoft Xbox", "Atari", "NEC", "SNK/Arcade"]
CATEGORY_BY_MANUFACTURER = {
    "Nintendo": "Nintendo", "Sony": "Sony PlayStation", "Sega": "Sega",
    "Microsoft": "Microsoft Xbox", "Atari": "Atari", "NEC": "NEC", "SNK / Arcade": "SNK/Arcade",
}


def import_workbook(workbook_path: Path, output_path: Path, enrichment_path: Path) -> int:
    try:
        import openpyxl
    except ImportError as exc:
        raise SystemExit("Für den Excel-Import bitte openpyxl installieren: pip install openpyxl") from exc
    enrichment = json.loads(enrichment_path.read_text(encoding="utf-8"))
    book = openpyxl.load_workbook(workbook_path, data_only=True)
    if "Emulatoren" not in book.sheetnames:
        raise SystemExit("Die Arbeitsmappe enthält kein Blatt 'Emulatoren'.")
    sheet = book["Emulatoren"]
    headers = [str(cell.value or "").strip() for cell in sheet[1]]
    required = ["Hersteller", "Konsole", "Bester Emulator", "Plattformen", "PC-Anforderung", "Offizielle Downloadquelle"]
    if headers[:6] != required or not headers[6].startswith("Download-Anleitung") or headers[7] != "Wichtiger Hinweis":
        raise SystemExit("Die Spalten des Blatts Emulatoren stimmen nicht mit dem erwarteten Format überein.")

    items = []
    official_sources = set(enrichment["official_sources"])
    matched = set()
    for row in sheet.iter_rows(min_row=2):
        values = [str(cell.value or "").strip() for cell in row[:8]]
        manufacturer, console, emulator, platforms, requirement, source_label, instructions, _old_note = values
        # Empty rows and the trailing table note are not emulator entries.
        if not manufacturer or not console:
            continue
        key = f"{manufacturer}|{console}"
        override = enrichment["entries"].get(key)
        if override is None:
            raise SystemExit(f"Neue Tabellenzeile {row[0].row}: Bitte die überprüfte Quelle und Anleitung für '{key}' in {enrichment_path.name} ergänzen.")
        if key in matched:
            raise SystemExit(f"Die Tabellenzeile '{key}' ist doppelt vorhanden.")
        matched.add(key)
        source_url = row[5].hyperlink.target if row[5].hyperlink else override.get("official_url")
        if not source_url or urlsplit(source_url).scheme != "https":
            raise SystemExit(f"Tabellenzeile {row[0].row}: Eine offizielle HTTPS-Quelle fehlt.")
        item = {
            "id": override["id"], "kategorie": CATEGORY_BY_MANUFACTURER.get(manufacturer, manufacturer),
            "hersteller": manufacturer, "konsole": console,
            "emulator": override.get("emulator", emulator), "plattformen": platforms,
            "pc_anforderung": requirement, "official_url": override.get("official_url", source_url),
            "download_anleitung": override.get("download_anleitung", instructions),
            "hinweis": f"{LEGAL_NOTICE} {override['hinweis']}",
            "install_methode": override.get("install_methode", "manuell"),
            "download_muster": override.get("download_muster", ""), "exe": override.get("exe", ""),
            "manuelle_schritte": override["manuelle_schritte"],
            "install_begruendung": override["install_begruendung"],
            "quelle_zeile": row[0].row, "quelle_beschreibung": source_label,
            "quelle_emulator": emulator, "verified_at": CHECKED,
        }
        for field in (
            "github_repo", "archive_type", "alternative_exes", "checksum_muster", "direct_url",
            "verified_release", "verified_asset", "verified_asset_sha256", "auto_emulator", "todo",
            "direct_version", "sha256", "checksum_url", "winget_id", "winget_version", "winget_manifest_url",
            "launch_args", "launch_profiles", "launch_extensions", "launch_note", "bios", "bios_note", "backup_paths", "backup_note",
            "controller_config", "controller_adapter", "controller_manual", "controller_note",
        ):
            if field in override:
                item[field] = override[field]
        items.append(item)
    unused = set(enrichment["entries"]) - matched
    if unused:
        raise SystemExit(f"Überprüfte Tabellenzeilen fehlen in der Arbeitsmappe: {', '.join(sorted(unused))}")
    result = {
        "schema_version": 1, "categories": CATEGORIES,
        "official_sources": sorted(official_sources),
        "source_workbook": (workbook_path.resolve().relative_to(ROOT).as_posix()
                            if workbook_path.resolve().is_relative_to(ROOT) else workbook_path.name),
        "source_sheet": "Emulatoren", "verified_at": CHECKED,
        "emulators": items,
    }
    output_path.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    Catalog(output_path)  # Validate the generated file with exactly the runtime rules.
    book.close()
    return len(items)


def main() -> None:
    parser = argparse.ArgumentParser(description="Den geprüften Emulator-Katalog aus Excel erzeugen.")
    parser.add_argument("workbook", nargs="?", type=Path, default=ROOT / "docs" / "emulatoren_mit_downloadanleitungen.xlsx")
    parser.add_argument("--output", type=Path, default=ROOT / "catalog.json")
    parser.add_argument("--enrichment", type=Path, default=ROOT / "tools" / "catalog_enrichment.json")
    args = parser.parse_args()
    count = import_workbook(args.workbook, args.output, args.enrichment)
    print(f"{count} Emulatoren nach {args.output} importiert und validiert.")


if __name__ == "__main__":
    main()
